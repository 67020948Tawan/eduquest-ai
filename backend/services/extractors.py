# ==============================================================================
# services/extractors.py — Document Text Extraction Layer
# สกัดข้อความจากไฟล์บทเรียน (PDF / DOCX / PPTX / TXT) ให้เป็น Plain Text
#
# รองรับ PDF ทั้งแบบ:
#   1. Text-based   → สกัดข้อความตรงด้วย PyMuPDF (เร็ว)
#   2. Image-based  → OCR fallback ด้วย Tesseract (tha+eng) เมื่อไม่มีข้อความฝัง
# ==============================================================================

import os
import shutil

TESSERACT_PATH_WIN = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
TESSDATA_DIR = os.path.join(os.path.dirname(__file__), "..", "tessdata")


def _resolve_tesseract_cmd() -> str | None:
    """หา binary ของ Tesseract ข้ามแพลตฟอร์ม: env > PATH > Windows default"""
    return os.getenv("TESSERACT_CMD") or shutil.which("tesseract") or TESSERACT_PATH_WIN

MIN_TEXT_LENGTH = 20
OCR_MAX_PAGES = 40  # จำกัดจำนวนหน้า OCR กัน request ยาวเกินไป


class ExtractionError(Exception):
    """Raised เมื่อสกัดข้อความจากไฟล์ไม่สำเร็จ"""


def _extract_pdf(file_path: str) -> str:
    """ดึงข้อความที่ฝังใน PDF ทีละหน้า + ใส่ tag [Page N] ให้ AI อ้างอิงตำแหน่งได้"""
    import fitz  # PyMuPDF

    parts: list[str] = []
    with fitz.open(file_path) as doc:
        for page_number, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                parts.append(f"[Page {page_number}]\n{text}")
    return "\n\n".join(parts)


def _ocr_pdf(file_path: str) -> str:
    """OCR fallback: render หน้ากระดาษเป็นภาพแล้วอ่านด้วย Tesseract (ไทย+อังกฤษ)"""
    import io

    import fitz

    cmd = _resolve_tesseract_cmd()
    if not cmd or not os.path.exists(cmd):
        raise ExtractionError(
            "ไฟล์ PDF นี้เป็นภาพสแกน (ไม่มีข้อความฝัง) และยังไม่ได้ติดตั้ง Tesseract OCR"
        )

    import pytesseract
    from PIL import Image

    pytesseract.pytesseract.tesseract_cmd = cmd
    # ใช้ tessdata ที่มากับโปรเจกต์ถ้ามี (local dev) — ไม่งั้นปล่อยให้ใช้ของระบบ (Docker)
    if "TESSDATA_PREFIX" not in os.environ:
        bundled = os.path.abspath(TESSDATA_DIR)
        if os.path.exists(os.path.join(bundled, "tha.traineddata")):
            os.environ["TESSDATA_PREFIX"] = bundled

    parts: list[str] = []
    with fitz.open(file_path) as doc:
        pages = list(doc)[:OCR_MAX_PAGES]
        for page_number, page in enumerate(pages, start=1):
            # render 200 DPI — สมดุลระหว่างความคมกับความเร็ว
            pix = page.get_pixmap(dpi=200)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            text = pytesseract.image_to_string(image, lang="tha+eng").strip()
            if text:
                parts.append(f"[Page {page_number}]\n{text}")

    return "\n\n".join(parts)


def _extract_docx(file_path: str) -> str:
    """ดึงข้อความทุกย่อหน้าใน Word (.docx) — ข้ามย่อหน้าว่าง"""
    from docx import Document as DocxDocument

    doc = DocxDocument(file_path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_pptx(file_path: str) -> str:
    """ดึงข้อความทุก text frame ใน PowerPoint (.pptx) พร้อม tag [Slide N]"""
    from pptx import Presentation

    prs = Presentation(file_path)
    parts: list[str] = []
    for slide_number, slide in enumerate(prs.slides, start=1):
        texts = [shape.text.strip() for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
        if texts:
            parts.append(f"[Slide {slide_number}]\n" + "\n".join(texts))
    return "\n\n".join(parts)


def _extract_plain(file_path: str) -> str:
    """อ่านไฟล์ข้อความล้วน (.txt/.md) — ข้าม byte ที่ decode ไม่ได้"""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


_EXTRACTORS = {
    ".pdf": _extract_pdf,
    ".docx": _extract_docx,
    ".pptx": _extract_pptx,
}

ALLOWED_EXTENSIONS = set(_EXTRACTORS.keys()) | {".txt", ".md"}


def extract_text(file_path: str, mime_type: str | None = None) -> str:
    """
    สกัดข้อความจากไฟล์ตามนามสกุล
    - PDF: ถ้าไม่มีข้อความฝัง (scanned/image-based) → OCR อัตโนมัติ
    """
    if not os.path.exists(file_path):
        raise ExtractionError("ไม่พบไฟล์บนเซิร์ฟเวอร์")

    extension = os.path.splitext(file_path)[1].lower()
    if not extension and mime_type:
        mime_map = {
            "application/pdf": ".pdf",
            "text/plain": ".txt",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation": ".pptx",
        }
        extension = mime_map.get(mime_type, "")

    if extension in _EXTRACTORS:
        extractor = _EXTRACTORS[extension]
    elif extension in (".txt", ".md"):
        extractor = _extract_plain
    else:
        raise ExtractionError(
            f"ไม่รองรับไฟล์นามสกุล '{extension or 'ไม่ทราบ'}' "
            f"(รองรับ: {', '.join(sorted(ALLOWED_EXTENSIONS))})"
        )

    try:
        text = (extractor(file_path) or "").strip()

        # PDF สแกน → ข้อความน้อยกว่า threshold → ลอง OCR
        if extension == ".pdf" and len(text) < MIN_TEXT_LENGTH:
            print(f"🔍 [Extractor] No embedded text — falling back to OCR: {file_path}")
            text = (_ocr_pdf(file_path) or "").strip()
    except ImportError as e:
        raise ExtractionError(f"ยังไม่ได้ติดตั้ง library สำหรับไฟล์ประเภทนี้: {e}") from e
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"อ่านไฟล์ไม่สำเร็จ: {e}") from e

    if len(text.strip()) < MIN_TEXT_LENGTH:
        raise ExtractionError(
            "สกัดข้อความจากไฟล์ไม่ได้ (ไฟล์อาจเป็นภาพล้วน/สแกนเบลอ หรือเนื้อหาว่างเปล่า)"
        )
    return text.strip()
